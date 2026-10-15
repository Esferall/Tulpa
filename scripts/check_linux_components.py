"""Run installed Linux media components on synthetic local files, without accounts."""
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import wave

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'scripts')]


def documents(folder):
    from parse_artifact import parse
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    expected='Tulpa Linux 2026'
    doc=Document();doc.add_paragraph(expected);doc.save(folder/'sample.docx')
    ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text=expected
    ppt.save(folder/'sample.pptx')
    book=Workbook();book.active['A1']=expected;book.save(folder/'sample.xlsx')
    writer=PdfWriter();page=writer.add_blank_page(width=600,height=800)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),
                          NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 18 Tf 40 700 Td (Tulpa Linux 2026) Tj ET')
    page[NameObject('/Contents')]=stream
    with (folder/'sample.pdf').open('wb') as out:writer.write(out)
    for extension in ('docx','pptx','xlsx','pdf'):
        result=parse(folder/('sample.'+extension),extension)
        assert expected in '\n'.join(c['text'] for c in result['chunks']),(extension,result)
        print('PASS: native '+extension+' parser',flush=True)


def ocr():
    from PIL import Image,ImageDraw,ImageFont
    from chatlocal.vision import VisionProvider
    image=Image.new('RGB',(560,110),'white');draw=ImageDraw.Draw(image)
    draw.text((20,20),'HELLO 2026',fill='black',font=ImageFont.load_default(size=48))
    buffer=BytesIO();image.save(buffer,format='PNG')
    result=VisionProvider({})._ocr([(0,buffer.getvalue())])
    assert '2026' in result['ocr'],result
    print('PASS: offline OCR inference',flush=True)


def voice(folder):
    import imageio_ffmpeg
    from chatlocal.platform_process import owned_popen,terminate_group
    from chatlocal.voice import config
    cfg=config()
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),'-version'],stdout=subprocess.DEVNULL,check=True,timeout=15)
    wav=folder/'silence.wav'
    with wave.open(str(wav),'wb') as output:
        output.setnchannels(1);output.setsampwidth(2);output.setframerate(16000)
        output.writeframes(bytes(16000*2*2))
    destination=folder/'transcript'
    args=[str(cfg['executable']),'-m',str(cfg['model']),'-f',str(wav),'-l','en','-t','2',
          '-oj','-of',str(destination),'-ng','-np']
    process=owned_popen(args,memory_bytes=2*1024**3,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
    try:
        _,error=process.communicate(timeout=180)
        assert process.returncode==0,error.decode(errors='replace')[-3000:]
        data=json.loads(destination.with_suffix('.json').read_text())
        assert 'transcription' in data and 'result' in data
    finally:terminate_group(process)
    print('PASS: native ffmpeg and whisper.cpp model inference (synthetic silence)',flush=True)


if __name__=='__main__':
    with tempfile.TemporaryDirectory(dir=ROOT/'.tmp',prefix='components-check-') as temporary:
        folder=Path(temporary)
        documents(folder);ocr();voice(folder)
